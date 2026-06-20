import os
import pandas as pd
import numpy as np
from prophet import Prophet
import itertools
import warnings
import logging
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Gereksiz uyarıları gizle (Temiz terminal çıktısı için)
warnings.filterwarnings('ignore')
logging.getLogger('cmdstanpy').setLevel(logging.WARNING)

print("🚀 1. Veri Setleri Okunuyor...")

def load_data(base_name):
    if os.path.exists(base_name + '.csv'): 
        return pd.read_csv(base_name + '.csv')
    elif os.path.exists(base_name + '.xlsx'): 
        return pd.read_excel(base_name + '.xlsx')
    else: 
        raise FileNotFoundError(f"HATA: '{base_name}' adlı dosya bulunamadı! Lütfen dosya isimlerini kontrol edin.")

# Verileri yükleme
df_demand = load_data('talep') 
df_coords = load_data('koordinatlar')
df_rentals = load_data('kiralik')
df_capacity = load_data('kapasite')

# Sütun başlıklarındaki olası görünmez boşlukları temizleme
df_demand.columns = df_demand.columns.str.strip()
df_coords.columns = df_coords.columns.str.strip()
df_rentals.columns = df_rentals.columns.str.strip()
df_capacity.columns = df_capacity.columns.str.strip()

# Kuş Uçuşu (Haversine) Mesafe Algoritması
def calculate_haversine(lat1, lon1, lat2, lon2):
    R = 6371.0
    lat1, lon1, lat2, lon2 = map(np.radians, [lat1, lon1, lat2, lon2])
    d = 2 * np.arctan2(np.sqrt(np.sin((lat2 - lat1)/2)**2 + np.cos(lat1) * np.cos(lat2) * np.sin((lon2 - lon1)/2)**2), 
                       np.sqrt(1 - (np.sin((lat2 - lat1)/2)**2 + np.cos(lat1) * np.cos(lat2) * np.sin((lon2 - lon1)/2)**2)))
    return R * d

coords = df_coords.set_index('Transfer Merkezi').to_dict('index')

# Maliyet ve Kapasite Sözlükleri (Hızlı işlem için)
cap_dict = dict(zip(df_capacity['Araç Adı'], df_capacity['Kapasite (desi)']))
spot_fixed = dict(zip(df_capacity['Araç Adı'], df_capacity['Spot Araç Sabit Günlük Maliyet (TL)']))
spot_km = dict(zip(df_capacity['Araç Adı'], df_capacity['Spot Kilometre Başına Maliyet (TL)']))
rent_fixed = dict(zip(df_capacity['Araç Adı'], df_capacity['Kiralık Araç Günlük Kira (TL)']))
rent_km = dict(zip(df_capacity['Araç Adı'], df_capacity['Kiralık Araç Kilometre Başına Maliyet (TL)']))

vehicles = df_capacity.to_dict('records') 

# Jürinin İstediği " TM" Formatlayıcı
def format_tm(name):
    name_str = str(name).strip()
    return name_str if name_str.endswith(" TM") else f"{name_str} TM"

print("📈 2. Prophet Algoritması ile Talep Tahmini Yapılıyor...")
forecast_list = []
for (origin, dest), group in df_demand.groupby(['Çıkış Transfer Merkezi', 'Varış Transfer Merkezi']):
    df_p = group[['Tarih', 'Toplam Desi']].rename(columns={'Tarih': 'ds', 'Toplam Desi': 'y'})
    df_p['ds'] = pd.to_datetime(df_p['ds'])
    
    # Makine Öğrenmesi Modeli
    m = Prophet(yearly_seasonality=False, daily_seasonality=True)
    m.fit(df_p)
    future = pd.DataFrame({'ds': pd.date_range(start='2026-05-11', end='2026-05-17')})
    fcst = m.predict(future)
    
    for _, row in fcst.iterrows():
        forecast_list.append({
            'Tarih': row['ds'].strftime('%Y-%m-%d'),
            'raw_orig': origin,
            'raw_dest': dest,
            'Tahmin Edilen Desi': max(0, round(row['yhat'], 2))
        })

df_forecast_raw = pd.DataFrame(forecast_list)

print("🚚 3. Araç Optimizasyonu (Spot %10 Kuralı ile) Yapılıyor...")
plan_results = []
total_system_cost = 0.0

rentals_grouped = df_rentals.groupby(['Çıkış Transfer Merkezi', 'Varış Transfer Merkezi', 'Araç Türü'])['Araç sayısı'].sum().reset_index()

for (orig, dest), group in df_forecast_raw.groupby(['raw_orig', 'raw_dest']):
    dist = calculate_haversine(coords[orig]['Enlem'], coords[orig]['Boylam'], coords[dest]['Enlem'], coords[dest]['Boylam']) if orig in coords and dest in coords else 0
    
    r_here = rentals_grouped[(rentals_grouped['Çıkış Transfer Merkezi'] == orig) & (rentals_grouped['Varış Transfer Merkezi'] == dest)]
    rental_fleet = []
    for _, r in r_here.iterrows():
        rental_fleet.extend([r['Araç Türü']] * r['Araç sayısı'])
            
    rollover = 0
    for _, row in group.sort_values('Tarih').iterrows():
        date = row['Tarih']
        daily_demand = row['Tahmin Edilen Desi'] + rollover
        current_demand = daily_demand
        
        # 1. Öncelikli Kiralık Araç Ataması (Sunk Cost Kuralı)
        for v_type in rental_fleet:
            cap = cap_dict[v_type]
            allocated = min(current_demand, cap)
            current_demand -= allocated
            cost = rent_fixed[v_type] + dist * rent_km[v_type]
            total_system_cost += cost
            
            v_name = v_type if "Kiralık" in v_type else f"Kiralık {v_type}"
            plan_results.append({
                'Tarih': date, 
                'Araç Tipi': v_name, 
                'Çıkış TM': format_tm(orig), 
                'Varış TM': format_tm(dest),
                'Atanan Desi': round(allocated, 2),
                'Maliyet': round(cost, 2)
            })
            
        # 2. Spot Araç Ataması (%10 Minimum Doluluk ve Devir Kısıtı)
        if current_demand > 0:
            best_cost, best_combo, best_carried = float('inf'), None, 0
            ranges = [range(int(current_demand // v['Kapasite (desi)']) + 3) for v in vehicles]
            
            for combo in itertools.product(*ranges):
                cap_total = sum(c * v['Kapasite (desi)'] for c, v in zip(combo, vehicles))
                min_10_percent = sum(c * 0.1 * v['Kapasite (desi)'] for c, v in zip(combo, vehicles))
                cost = sum(c * (v['Spot Araç Sabit Günlük Maliyet (TL)'] + dist * v['Spot Kilometre Başına Maliyet (TL)']) for c, v in zip(combo, vehicles))
                d_carried = min(current_demand, cap_total)
                
                # Jürinin Acımasız Kısıtı
                if d_carried >= min_10_percent and (current_demand - d_carried) < 560:
                    if cost < best_cost:
                        best_cost, best_combo, best_carried = cost, combo, d_carried
                        
            if best_combo:
                for idx, count in enumerate(best_combo):
                    for _ in range(count):
                        v_type = vehicles[idx]['Araç Adı']
                        cap = vehicles[idx]['Kapasite (desi)']
                        allocated = min(current_demand, cap)
                        current_demand -= allocated
                        
                        s_cost = spot_fixed[v_type] + dist * spot_km[v_type]
                        total_system_cost += s_cost
                        
                        v_name = v_type if "Spot" in v_type else f"Spot {v_type}"
                        plan_results.append({
                            'Tarih': date, 
                            'Araç Tipi': v_name, 
                            'Çıkış TM': format_tm(orig), 
                            'Varış TM': format_tm(dest),
                            'Atanan Desi': round(allocated, 2),
                            'Maliyet': round(s_cost, 2)
                        })
                rollover = current_demand
            else:
                rollover = current_demand
        else:
            rollover = 0

df_plan = pd.DataFrame(plan_results)

# Q&A Formatına Tam Uyumlu Sütun Sıralaması
df_forecast = df_forecast_raw.copy()
df_forecast['Çıkış TM'] = df_forecast['raw_orig'].apply(format_tm)
df_forecast['Varış TM'] = df_forecast['raw_dest'].apply(format_tm)
df_forecast = df_forecast[['Tarih', 'Çıkış TM', 'Varış TM', 'Tahmin Edilen Desi']]

print("💾 4. Excel Dosyaları Şık Tasarımla (.xlsx formatında) Oluşturuluyor...")

# Excel Tasarım Fonksiyonu
def format_excel_and_save(df, filename, sheet_name="Sayfa1"):
    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]
        
        # Mavi Arka Plan, Beyaz Kalın Yazı
        header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        
        for col_idx, col_name in enumerate(df.columns, 1):
            col_letter = get_column_letter(col_idx)
            header_cell = worksheet.cell(row=1, column=col_idx)
            header_cell.fill = header_fill
            header_cell.font = header_font
            header_cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # Sütun Genişliklerini Otomatik Ayarlama 
            max_len = len(str(col_name))
            for row_val in df[col_name]:
                if pd.notna(row_val):
                    max_len = max(max_len, len(str(row_val)))
            worksheet.column_dimensions[col_letter].width = max_len + 5
            
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=len(df.columns)):
            for cell in row:
                cell.alignment = Alignment(horizontal='center', vertical='center')

# Çıktıları Kaydet
format_excel_and_save(df_forecast, 'Tahminlenen_Talep_Ciktisi.xlsx')
format_excel_and_save(df_plan, 'Arac_Planlama_Ciktisi.xlsx')

print("\n" + "="*50)
print("✅ İŞLEM TAMAMLANDI!")
print(f"💰 TOPLAM SİSTEM MALİYETİ: {total_system_cost:,.2f} TL")
print("="*50)
